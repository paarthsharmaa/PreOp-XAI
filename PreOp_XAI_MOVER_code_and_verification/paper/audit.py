# -*- coding: utf-8 -*-
"""Independent audit of the built paper against the exported MOVER results. Does NOT import content.py: every
expected value is recomputed here from the raw JSON files (main runs, 200 repeated-split files, cohort files) and
looked up in the .docx text; qualitative claims are re-derived and checked; structure is checked.
usage: python3 audit.py out/PreOp_XAI_MOVER_IEEE.docx [results_dir]"""
import glob
import hashlib
import json
import os
import re
import sys
import zipfile

import numpy as np
from docx import Document

DOCX = sys.argv[1] if len(sys.argv) > 1 else "out/PreOp_XAI_MOVER_IEEE.docx"
R = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results")
J = lambda *p: json.load(open(os.path.join(R, *p)))
doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]
cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
TEXT = "\n".join(paras + cells).replace("−", "-")
BODY = "\n".join(paras).replace("−", "-")
results = []


def check(name, ok, detail=""):
    results.append((bool(ok), name, detail))


def has(s, name=None):
    s = s.replace("−", "-")
    check(name or f"text contains '{s}'", s in TEXT, "" if s in TEXT else "MISSING")


f3 = lambda x: f"{x:.3f}"
f2 = lambda x: f"{x:.2f}"
f4 = lambda x: f"{x:.4f}"


def iv(d, k, fmt=f3):
    lo, hi = fmt(d["intervals"][k]["lower"]), fmt(d["intervals"][k]["upper"])
    return f"{lo} to {hi}" if (lo.startswith("-") or hi.startswith("-")) else f"{lo}–{hi}"


M, I = J("main_mortality", "study.json"), J("main_icu", "study.json")
S = {"mortality": M, "icu": I}
L0 = {"mortality": J("lag0_mortality", "study.json"), "icu": J("lag0_icu", "study.json")}
FLOW, LABS, LABS0 = J("cohort", "mover_cohort_flow.json"), J("cohort", "mover_feature_stats.json"), J("cohort", "mover_feature_stats_lag0h.json")
RESP = J("cohort", "mover_respiratory_feasibility.json")["modelled"]
SC, DET = J("split_check.json"), J("determinism_check.json")
REP = {e: sorted([json.load(open(p)) for p in glob.glob(os.path.join(R, f"repeats_{e}", "seed*.json"))], key=lambda r: r["seed"])
       for e in S}

# ------------------------------------------------------------ data integrity
check("determinism rerun identical", DET["identical_except_runtime"] and not DET["differences"])
check("SHAP self-test passed", "PASS" in open(os.path.join(R, "shap_selftest.txt")).read())
for e in S:
    check(f"{e}: 100 repeats with seeds 1..100", [r["seed"] for r in REP[e]] == list(range(1, 101)))
rr = J("rerun_mortality", "study.json")
check("rerun evaluation equals main", rr["evaluation"] == M["evaluation"] and rr["external_vitaldb"]["models"] == M["external_vitaldb"]["models"])
check("main runs used the 24-h table", M["table"] == "mover_analysis.csv" and I["table"] == "mover_analysis.csv")
check("lag-0 runs used the lag-0 table", all(L0[e]["table"] == "mover_analysis_lag0h.csv" for e in S))
env = J("environment.json")
check("table hashes match environment record", M["table_sha256"] == env["inputs"]["mover_analysis.csv"]
      and L0["icu"]["table_sha256"] == env["inputs"]["mover_analysis_lag0h.csv"])
check("repository commit 6baa665", env["repo_commit"].startswith("6baa665"))

# ------------------------------------------------------------ cohort flow arithmetic and text
ex = FLOW["exclusions"]
elig = FLOW["rows"] - FLOW["exact_duplicates_removed"] - ex["operation_conflict"] - ex["not_adult"] - ex["invalid_or_donor_asa"] - ex["invalid_anesthesia_boundary"]
check("flow: eligible - later = indexed", elig - ex["later_eligible_operation"] - ex["ambiguous_first_operation_time"] == FLOW["indexed_patients"])
check("flow: indexed - unknown = modelled", FLOW["indexed_patients"] - FLOW["unlinked"] - FLOW["unknown_disposition"] == FLOW["mortality_known"])
check("modelled n equals study n", M["n"] == FLOW["mortality_known"] == I["n"] and M["events"] == FLOW["mortality_events"]
      and I["events"] == FLOW["icu_flag_yes_among_known"])
for v in (FLOW["patients"], FLOW["operations"], FLOW["exact_duplicates_removed"], ex["invalid_anesthesia_boundary"],
          FLOW["indexed_patients"], FLOW["mortality_known"], FLOW["icu_flag_yes_among_known"]):
    has(f"{v:,}")
for v in (ex["operation_conflict"], ex["not_adult"], ex["invalid_or_donor_asa"]):
    has(f"({v} excluded)" if v == ex["operation_conflict"] else f"({v})")
has(f"{100 * FLOW['mortality_events'] / FLOW['mortality_known']:.2f}%")
has(f"{100 * FLOW['icu_flag_yes_among_known'] / FLOW['mortality_known']:.1f}%")
for k in ("positive_only", "contradictory", "negative_only"):
    has(f"{RESP[k]:,}", f"respiratory {k} count")
has(f"{LABS['candidate_rows']:,}"); has(f"{LABS['reasons']['eligible']:,}"); has(f"{LABS['reasons']['future_collection']:,}")
has(f"{LABS['unmapped_code_or_unit']:,}")
LK = ("hemoglobin", "platelets", "creatinine", "bun", "sodium", "potassium", "glucose")
has(f"{min(LABS['missing_pct'][k] for k in LK):.1f}–{max(LABS['missing_pct'][k] for k in LK):.1f}%", "24-h missingness range")
has(f"{min(LABS0['missing_pct'][k] for k in LK):.1f}–{max(LABS0['missing_pct'][k] for k in LK):.1f}%", "lag-0 missingness range")
P = M["partition_n"]
for k in ("train", "validation", "test"):
    has(f"{P[k]:,}")
check("partitions identical for both endpoints", M["partition_n"] == I["partition_n"])
check("test deaths identical in all repeats (stratified)", {r["test_events"] for r in REP["mortality"]} == {M["partition_events"]["test"]})
has(f"{M['partition_events']['test']} deaths")
check("primary ICU test count typical (between 10th and 90th percentile share)", 0.1 < SC["alt_share_le_primary"] < 0.9)
has(f"{SC['alt_test_icu_min']:,}–{SC['alt_test_icu_max']:,}")

# ------------------------------------------------------------ Table I values recomputed from characteristics
C = M["characteristics"]; CI = I["characteristics"]["event"]; CV = M["external_vitaldb"]["characteristics"]["all"]
for c in (C["all"], C["event"], C["no_event"], CI, CV):
    has(f"{c['age_median']:.0f} ({c['age_q1']:.0f}–{c['age_q3']:.0f})", "Table I age")
    has(f"{c['female_pct']:.1f}", "Table I female")
    has(f"{c['asa_pct']['3']:.1f} / {c['asa_pct']['4']:.1f} / {c['asa_pct']['5']:.1f}", "Table I ASA III-V")
    has(f"{c['creatinine']['median']:.2f} ({c['creatinine']['q1']:.2f}–{c['creatinine']['q3']:.2f})", "Table I creatinine")
check("characteristics 'all' n equals cohort", C["all"]["n"] == FLOW["mortality_known"] and CV["n"] == 6019)
check("claim: deaths more often ASA IV-V", C["event"]["asa_pct"]["4"] + C["event"]["asa_pct"]["5"] > C["no_event"]["asa_pct"]["4"] + C["no_event"]["asa_pct"]["5"])
check("claim: deaths less often missing labs", max(C["event"][k]["missing_pct"] for k in LK) < min(C["all"][k]["missing_pct"] for k in LK))
check("claim: VitalDB healthier (ASA I-II share higher)", CV["asa_pct"]["1"] + CV["asa_pct"]["2"] > C["all"]["asa_pct"]["1"] + C["all"]["asa_pct"]["2"])
check("claim: VitalDB almost completely tested (<10% missing)", max(CV[k]["missing_pct"] for k in LK) < 10)
check("claim: deaths older", C["event"]["age_median"] > C["no_event"]["age_median"])

# ------------------------------------------------------------ Table II / discrimination
for e, d in S.items():
    for m in ("lr", "xgboost"):
        r = d["evaluation"]["models"][m]
        has(f"{f3(r['raw']['auroc'])} ({iv(d['evaluation'], f'{m}.raw.auroc')})", f"{e} {m} test AUROC+CI")
        has(f"{f3(r['raw']['average_precision'])} ({iv(d['evaluation'], f'{m}.raw.average_precision')})", f"{e} {m} test AP+CI")
        x = d["external_vitaldb"]["models"][m]
        has(f"{f3(x['raw']['auroc'])} ({iv(d['external_vitaldb'], f'{m}.raw.auroc')})", f"{e} {m} external AUROC+CI")
        has(f3(L0[e]["evaluation"]["models"][m]["raw"]["auroc"]), f"{e} {m} lag-0 AUROC")
        # independent recomputation of the repeated-split median and percentiles
        a = np.array([r_["test"][m]["raw"]["auroc"] for r_ in REP[e]])
        has(f"{f3(np.median(a))} ({f3(np.quantile(a, .025))}–{f3(np.quantile(a, .975))})", f"{e} {m} repeats AUROC")
    dd = d["evaluation"]["models"]
    has(f"{f3(dd['xgboost']['raw']['auroc'] - dd['lr']['raw']['auroc'])} ({iv(d['evaluation'], 'paired_xgboost_minus_lr.raw.auroc')})", f"{e} paired ΔAUROC")
    diff = np.array([r_["test"]["xgboost"]["raw"]["auroc"] - r_["test"]["lr"]["raw"]["auroc"] for r_ in REP[e]])
    dap = np.array([r_["test"]["xgboost"]["raw"]["average_precision"] - r_["test"]["lr"]["raw"]["average_precision"] for r_ in REP[e]])
    has(f"{100 * np.mean(diff > 0):.0f}% of splits" if e == "mortality" else f"{100 * np.mean(diff > 0):.0f}% of splits", f"{e} share XGB ahead")
    if e == "mortality":
        has(f"between {f3(np.quantile(diff, .025))} and {f3(np.quantile(diff, .975))}", "mortality Δ 95% range")
        has(f"LR had the higher AP in {100 * np.mean(dap < 0):.0f}%", "LR higher AP share")
        bound = max(abs(np.quantile(diff, .025)), abs(np.quantile(diff, .975)))
        has(f"within {f3(bound)} in 95%", "abstract Δ bound")
        icu = np.array([r_["test"]["xgboost"]["raw"]["auroc"] - r_["test"]["lr"]["raw"]["auroc"] for r_ in REP["icu"]])
        check("abstract bound also covers ICU", max(abs(np.quantile(icu, .025)), abs(np.quantile(icu, .975))) <= bound)
    else:
        has(f"median of only {f4(np.median(diff))}", "ICU Δ median")
        has(f"median {f4(np.median(dap))}", "ICU ΔAP median")
xi = I["external_vitaldb"]
check("claim: LR better externally for ICU (interval excludes 0)", xi["intervals"]["paired_xgboost_minus_lr.raw.auroc"]["upper"] < 0)
l0 = L0["icu"]["evaluation"]
check("claim: lag-0 ICU XGB ahead (interval excludes 0)", l0["intervals"]["paired_xgboost_minus_lr.raw.auroc"]["lower"] > 0)
check("claim: lag-0 AUROCs higher than primary", all(L0[e]["evaluation"]["models"][m]["raw"]["auroc"] > S[e]["evaluation"]["models"][m]["raw"]["auroc"] for e in S for m in ("lr", "xgboost")))
check("claim: ICU XGB advantage reversed externally", xi["models"]["xgboost"]["raw"]["auroc"] < xi["models"]["lr"]["raw"]["auroc"])

# ------------------------------------------------------------ calibration and thresholds
for e, d in S.items():
    for m in ("lr", "xgboost"):
        r, x = d["evaluation"]["models"][m], d["external_vitaldb"]["models"][m]
        has(f"{f2(r['calibration']['offset_intercept'])} / {f2(r['calibration']['joint_slope'])}", f"{e} {m} Table III MOVER")
        has(f"{f2(x['calibration']['offset_intercept'])} / {f2(x['calibration']['joint_slope'])}", f"{e} {m} Table III VitalDB")
        has(f"{f3(r['sensitivity'])} / {f3(r['specificity'])} / {f3(r['ppv'])}", f"{e} {m} Table III op MOVER")
        has(f"{f3(x['sensitivity'])} / {f3(x['specificity'])} / {f3(x['ppv'])}", f"{e} {m} Table III op VitalDB")
        check(f"{e} {m}: Brier below prevalence reference internally", r["raw"]["brier"] < r["training_prevalence_brier"]
              and r["calibrated"]["brier"] < r["training_prevalence_brier"])
        has(f2(d["selection"][m]["diagnostics"]["calibration_slope"]), f"{e} {m} validation slope")
    check(f"{e}: no non-positive validation slope over repeats",
          all(r_["selection"][m]["calibration_slope"] > 0 for r_ in REP[e] for m in ("lr", "xgboost")))
check("claim: mortality underestimated externally (CITL>0)", all(M["external_vitaldb"]["models"][m]["calibration"]["offset_intercept"] > 0 for m in ("lr", "xgboost")))
check("claim: ICU overestimated externally (CITL<0)", all(xi["models"][m]["calibration"]["offset_intercept"] < 0 for m in ("lr", "xgboost")))
check("claim: external ICU recalibrated Brier worse than reference", all(xi["models"][m]["calibrated"]["brier"] > xi["models"][m]["training_prevalence_brier"] for m in ("lr", "xgboost")))
mL = M["evaluation"]["models"]["lr"]; mX = M["evaluation"]["models"]["xgboost"]
has(f"LR flagged {mL['tp'] + mL['fp']} test patients"); has(f"detected {mL['tp']} of {M['partition_events']['test']} deaths")
has(f"XGBoost detected {mX['tp']} while flagging {mX['tp'] + mX['fp']}")
xm = M["external_vitaldb"]["models"]
has(f"detected only {xm['lr']['tp']} (LR) and {xm['xgboost']['tp']} (XGBoost) of {M['external_vitaldb']['events']} deaths")
sens = {m: np.mean([r_["test"][m]["sensitivity"] >= 0.8 for r_ in REP["mortality"]]) for m in ("lr", "xgboost")}
has(f"{100 * sens['lr']:.0f}%–{100 * sens['xgboost']:.0f}% of splits", "share sensitivity>=0.8")
cs = np.array([r_["test"]["lr"]["calibration"]["joint_slope"] for r_ in REP["mortality"]])
has(f"between {f2(np.quantile(cs, .025))} and {f2(np.quantile(cs, .975))} (LR)", "repeat recal. slope range")
has(f"{M['evaluation']['models']['lr']['missing_labs_mean']:.1f} of seven"); has(f"but {xm['lr']['missing_labs_mean']:.1f} in VitalDB")

# ------------------------------------------------------------ explanations and stability
CON = ["age", "sex", "asa", "hemoglobin", "platelets", "creatinine", "bun", "sodium", "potassium", "glucose"]
for e, d in S.items():
    for m in ("lr", "xgboost"):
        g = d["explain"]["models"][m]["global_mean_absolute_raw"]
        order = sorted(g, key=lambda k: -g[k])
        check(f"{e} {m}: ASA top in primary fit", order[0] == "asa")
        imp = np.array(d["figure_data"]["shap_refit_importance"][m])
        check(f"{e} {m}: ASA top in every refit", (imp.argmax(1) == CON.index("asa")).all())
        st = d["stability"]["models"][m]
        check(f"{e} {m}: 10 successful refits", st["status"] == "complete" and st["successful_refits"] == 10)
        has(f2(st["spearman_mean"]), f"{e} {m} Spearman"); has(f2(st["top5_jaccard_mean"]), f"{e} {m} Jaccard")
        # independent Spearman from the refit importance matrix (rank correlation of mean |SHAP| between refits)
        from scipy.stats import spearmanr
        vals = [spearmanr(imp[i], imp[j]).statistic for i in range(10) for j in range(i + 1, 10)]
        check(f"{e} {m}: Spearman recomputed from refit importances", abs(np.mean(vals) - st["spearman_mean"]) < 5e-3,
              f"{np.mean(vals):.4f} vs {st['spearman_mean']:.4f}")
        if e == "mortality" or m == "lr":
            check(f"{e} {m}: sex and age next after ASA in primary fit", set(order[1:3]) == {"sex", "age"}, str(order[:4]))
check("claim: ICU XGB primary fit has glucose second", sorted(I["explain"]["models"]["xgboost"]["global_mean_absolute_raw"].items(), key=lambda t: -t[1])[1][0] == "glucose")
check("claim: mortality rankings more stable than ICU", min(M["stability"]["models"][m]["spearman_mean"] for m in ("lr", "xgboost"))
      > max(I["stability"]["models"][m]["spearman_mean"] for m in ("lr", "xgboost")))
has(f"{100 * I['stability']['models']['xgboost']['top5_inclusion_frequency']['glucose']:.0f}% of XGBoost refits")
share = {(e, m): S[e]["explain"]["models"][m]["global_mean_absolute_raw"]["asa"] / sum(S[e]["explain"]["models"][m]["global_mean_absolute_raw"].values()) for e in S for m in ("lr", "xgboost")}
for e in S:
    lo, hi = sorted((share[(e, "lr")], share[(e, "xgboost")]))
    has(f"{100 * lo:.0f}%" if f"{100 * lo:.0f}" == f"{100 * hi:.0f}" else f"{100 * lo:.0f}%–{100 * hi:.0f}%", f"ASA share {e}")
has(f"{C['all']['asa_pct']['4'] + C['all']['asa_pct']['5']:.1f}% of the cohort", "ASA IV-V cohort share")
has(f"{C['event']['asa_pct']['4'] + C['event']['asa_pct']['5']:.1f}% of deaths", "ASA IV-V death share")

# ------------------------------------------------------------ structure
auth = [c.text for t in doc.tables[:2] for row in t.rows for c in row.cells]
names = [a.split("\n")[0] for a in auth]
check("five authors in order", names == ["Raghav Goel", "Bhavya Mansharamani", "Subhash Chand Gupta", "Prakhar Chandra", "Shipra Saraswat"], str(names))
check("every author: Department of CSE, Amity University", all("Department of CSE" in a and "Amity University" in a for a in auth))
check("no 'Amity School' anywhere", "Amity School" not in TEXT)
check("Prakhar and Shipra emails", "prakharchandra018@gmail.com" in auth[3] and "sshipra1510@gmail.com" in auth[4])
heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1") or p.style.name == "Heading1"]
check("sections", heads == ["Introduction", "Related Work", "Data and Methods", "Results", "Discussion", "Conclusion"], str(heads))
with zipfile.ZipFile(DOCX) as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    hashes = {hashlib.md5(z.read(n)).hexdigest() for n in media}
figs = sorted(glob.glob(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "figs", "fig[1-4]_*.png")))
check("4 figure files embedded byte-identical", len(figs) == 4 and {hashlib.md5(open(f, "rb").read()).hexdigest() for f in figs} <= hashes,
      f"{len(media)} media")
check("exactly 4 figures (images) in body", len(media) == 4, str(len(media)))
data_tables = [t for t in doc.tables[2:]]
check("exactly 3 data tables", len(data_tables) == 3, str(len(data_tables)))
for k in ("Fig. 1", "Fig. 2", "Fig. 3", "Fig. 4", "Table I", "Table II", "Table III"):
    check(f"{k} referenced in body text", k in BODY)
check("no Fig. 5 / Table IV references", "Fig. 5" not in BODY and "Table IV" not in BODY)
_body = "\n".join(paras[:paras.index("References")] if "References" in paras else paras)
_body = re.sub(r"\[(\d+)\]\s*[–-]\s*\[(\d+)\]", lambda m_: ", ".join(f"[{k}]" for k in range(int(m_.group(1)), int(m_.group(2)) + 1)), _body)
cites = [int(x) for x in re.findall(r"\[(\d+)\]", _body)]
first = []
for c_ in cites:
    if c_ not in first:
        first.append(c_)
nref = sum(1 for p in doc.paragraphs if p.style.name.lower().startswith("references"))
check("references cited in numerical order of first appearance", first == list(range(1, len(first) + 1)), str(first[:12]))
check("every reference cited", len(first) == nref, f"{len(first)} cited, {nref} listed")
check("reference count 28", nref == 28, str(nref))
words = sum(len(p.split()) for p in paras[paras.index("Introduction") if "Introduction" in paras else 0: paras.index("References") if "References" in paras else None])
check("body length 2,300-3,300 words", 2300 <= words <= 3300, str(words))
check("no leftover template text", not any(s in TEXT for s in ("Heading 1", "use style", "Given Name", "dept. name")))

bad = [r for r in results if not r[0]]
for ok, name, det in results:
    print(("PASS " if ok else "FAIL ") + name + (f"  [{det}]" if det else ""))
print(f"\n{len(results) - len(bad)}/{len(results)} checks passed")
if any("[email]" in a for a in auth):
    print("NOTE: email placeholders remain for", [a.split(chr(10))[0] for a in auth if "[email]" in a])
sys.exit(1 if bad else 0)
