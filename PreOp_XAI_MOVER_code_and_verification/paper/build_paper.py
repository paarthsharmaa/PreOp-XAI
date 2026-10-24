# -*- coding: utf-8 -*-
"""Build the PreOp-XAI paper on the (converted) IEEE conference template, using the template's own styles:
paper title, Author, Abstract, Keywords, Heading 1/2/5, Body Text, equation, table head, table col head, table copy,
figure caption, references (auto-numbered by the template's numbering definitions)."""
import copy, re, sys, os, json, hashlib, subprocess, tempfile, zipfile
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, Inches, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
import content as C
from refs import REFS

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
PREVIEW = {"fn": None}
BODY_W = 11906 - 2 * 907                 # twips, body text width (template body section margins)
COL_W = (BODY_W - 360) // 2              # 4866 twips
TNR = "Times New Roman"

# ------------------------------------------------------------------ citations (numbered by first appearance)
order = []
GROUP = re.compile(r"\[\[[a-z_0-9]+\]\](?:,\s*\[\[[a-z_0-9]+\]\])*")


def cite(text):
    def rep(m):
        keys = re.findall(r"\[\[([a-z_0-9]+)\]\]", m.group(0)); nums = []
        for k in keys:
            assert k in REFS, k
            if k not in order:
                order.append(k)
            nums.append(order.index(k) + 1)
        nums = sorted(set(nums)); out = []; i = 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else ", ".join(f"[{x}]" for x in nums[i:j + 1]))
            i = j + 1
        return ", ".join(out)
    return GROUP.sub(rep, text)


# ------------------------------------------------------------------ math
def latex_to_omml(latex_list):
    md = "\n\n".join(f"X ${l}$" for l in latex_list)
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "m.md"); out = os.path.join(td, "m.docx")
        open(src, "w").write(md)
        subprocess.run(["pandoc", src, "-o", out], check=True)
        xml = zipfile.ZipFile(out).read("word/document.xml")
    maths = etree.fromstring(xml).findall(".//{%s}oMath" % M_NS)
    assert len(maths) == len(latex_list), (len(maths), len(latex_list))
    return dict(zip(latex_list, maths))


def sized_math(el, size_pt):
    el = copy.deepcopy(el)
    for r in el.iter("{%s}r" % M_NS):
        for old in r.findall("{%s}rPr" % W_NS):
            r.remove(old)
        rpr = etree.Element("{%s}rPr" % W_NS)
        f = etree.SubElement(rpr, "{%s}rFonts" % W_NS)
        for a in ("ascii", "hAnsi", "cs"):
            f.set("{%s}%s" % (W_NS, a), "Cambria Math")
        for tag in ("sz", "szCs"):
            e = etree.SubElement(rpr, "{%s}%s" % (W_NS, tag)); e.set("{%s}val" % W_NS, str(int(round(size_pt * 2))))
        mrpr = r.find("{%s}rPr" % M_NS)
        r.insert(list(r).index(mrpr) + 1 if mrpr is not None else 0, rpr)
    return el


MATH = {}


def collect_math():
    items = set()
    def scan(t):
        for m in re.findall(r"\$([^$]+)\$", t):
            items.add(m)
    for h1, h2, paras in C.SECTIONS:
        for p in paras:
            if p.startswith("EQ:"):
                items.add(p[3:])
            elif not p.startswith(("FIG:", "TAB:")):
                scan(p)
    for c in C.CAPTIONS.values():
        scan(c)
    for t in C.TABLES.values():
        for cell in t["header"] + [x for r in t["rows"] for x in r]:
            scan(cell)
    items = sorted(items)
    MATH.update(latex_to_omml(items))


# ------------------------------------------------------------------ low-level XML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def new_p(style=None):
    p = el("w:p"); ppr = el("w:pPr"); p.append(ppr)
    if style:
        ppr.append(el("w:pStyle", **{"w:val": style}))
    return p


def ppr_add(p, child):
    """Insert a pPr child respecting the schema order used here (pStyle first, sectPr last)."""
    ppr = p.find(qn("w:pPr"))
    order_ = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:widowControl", "w:numPr", "w:tabs",
              "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr"]
    tag = child.tag.split("}")[1]
    idx = order_.index("w:" + tag)
    pos = len(ppr)
    for i, c in enumerate(ppr):
        ct = "w:" + c.tag.split("}")[1]
        if ct in order_ and order_.index(ct) > idx:
            pos = i; break
    ppr.insert(pos, child)
    return child


def run(text, size=None, bold=False, italic=False, sup=False, sub=False, smallcaps=False, font=None):
    r = el("w:r"); rpr = el("w:rPr"); r.append(rpr)
    if font:
        rpr.append(el("w:rFonts", **{"w:ascii": font, "w:hAnsi": font, "w:cs": font}))
    if bold: rpr.append(el("w:b"))
    if italic: rpr.append(el("w:i"))
    if smallcaps: rpr.append(el("w:smallCaps"))
    if size:
        rpr.append(el("w:sz", **{"w:val": int(round(size * 2))})); rpr.append(el("w:szCs", **{"w:val": int(round(size * 2))}))
    if sup or sub:
        rpr.append(el("w:vertAlign", **{"w:val": "superscript" if sup else "subscript"}))
    t = el("w:t"); t.text = text; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); r.append(t)
    return r


def br(kind=None):
    r = el("w:r"); b = el("w:br")
    if kind: b.set(qn("w:type"), kind)
    r.append(b); return r


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\$[^$]+\$)")


def rich(p, text, size=None, italic=False, bold=False, font=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            p.append(run(tok[2:-2], size, bold=True, italic=italic, font=font))
        elif tok.startswith("*"):
            p.append(run(tok[1:-1], size, bold=bold, italic=not italic, font=font))
        elif tok.startswith("$"):
            p.append(math_inline(tok[1:-1], size or 10))
        else:
            p.append(run(tok, size, bold=bold, italic=italic, font=font))
    return p


def math_inline(latex, size):
    if PREVIEW["fn"]:
        return picture_run(*PREVIEW["fn"](latex, size))
    return sized_math(MATH[latex], size)


# pictures: build them through python-docx so relationships are handled
DOC = None


def picture_run(path, height_pt=None, width_in=None):
    tmp = DOC.add_paragraph()
    r = tmp.add_run()
    if width_in:
        r.add_picture(path, width=Inches(width_in))
    else:
        r.add_picture(path, height=Pt(height_pt))
    r_el = r._r
    tmp._p.getparent().remove(tmp._p)
    return r_el


# ------------------------------------------------------------------ template surgery
def build(path):
    global DOC
    order.clear()
    doc = Document("template_T.docx"); DOC = doc
    body = doc.element.body
    kids = list(body)
    title_p, title_end = kids[0], kids[3]
    body_end = kids[90]              # paragraph carrying the 2-column body sectPr
    final_sect = kids[92]
    for i, k in enumerate(kids):
        if i not in (0, 3, 90, 92):
            body.remove(k)
    # clear the red guidance runs in the body-end paragraph; keep its sectPr
    for r in body_end.findall(qn("w:r")):
        body_end.remove(r)
    body_sect = body_end.find(qn("w:pPr")).find(qn("w:sectPr"))
    bst = body_end.find(qn("w:pPr")).find(qn("w:pStyle"))
    bst.set(qn("w:val"), "Normal")                       # was 'references' (auto-numbered) -> would print "[31]"
    for tag in ("w:spacing", "w:jc", "w:numPr", "w:ind"):
        e = body_end.find(qn("w:pPr")).find(qn(tag))
        if e is not None:
            body_end.find(qn("w:pPr")).remove(e)
    ppr_add(body_end, el("w:spacing", **{"w:before": 0, "w:after": 0, "w:line": 20, "w:lineRule": "exact"}))
    # paragraph for the (empty) final single-column section
    tail = new_p(); body_end.addnext(tail)
    ppr_add(tail, el("w:spacing", **{"w:before": 0, "w:after": 0, "w:line": 20, "w:lineRule": "exact"}))
    trpr = el("w:rPr"); trpr.append(el("w:sz", **{"w:val": 2})); ppr_add(tail, trpr)

    def insert(elm):
        body_end.addprevious(elm); return elm

    # ---- title
    for r in title_p.findall(qn("w:r")):
        title_p.remove(r)
    title_p.append(run(C.TITLE))

    # ---- authors: borderless tables inside the single-column title section (3 + 2, centred)
    def author_table(auths, width_tw):
        t = el("w:tbl"); tp = el("w:tblPr"); t.append(tp)
        tp.append(el("w:tblW", **{"w:w": width_tw, "w:type": "dxa"}))
        tp.append(el("w:jc", **{"w:val": "center"}))
        b = el("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b.append(el(f"w:{e}", **{"w:val": "nil"}))
        tp.append(b)
        tp.append(el("w:tblLayout", **{"w:type": "fixed"}))
        m = el("w:tblCellMar")
        for k in ("top", "left", "bottom", "right"):
            m.append(el(f"w:{k}", **{"w:w": 0 if k in ("top", "bottom") else 60, "w:type": "dxa"}))
        tp.append(m)
        g = el("w:tblGrid"); t.append(g)
        cw = width_tw // len(auths)
        for _ in auths:
            g.append(el("w:gridCol", **{"w:w": cw}))
        tr = el("w:tr"); t.append(tr)
        trpr = el("w:trPr"); trpr.append(el("w:cantSplit")); tr.append(trpr)
        for a in auths:
            tc = el("w:tc"); tr.append(tc)
            tcp = el("w:tcPr"); tcp.append(el("w:tcW", **{"w:w": cw, "w:type": "dxa"})); tc.append(tcp)
            p = new_p("Author")
            ppr_add(p, el("w:spacing", **{"w:before": 0, "w:after": 0}))
            tc.append(p)
            lines = [(a["name"], False), (a["dept"], True), (a["org"], True), (a["city"], False), (a["email"], False)]
            for i, (txt, it) in enumerate(lines):
                if i:
                    p.append(br())
                p.append(run(txt, 10 if i == 0 else 9, italic=it))
        return t

    tA = author_table(C.AUTHORS[:3], 10120)
    title_end.addprevious(tA)
    if len(C.AUTHORS) > 3:
        tB = author_table(C.AUTHORS[3:], 10120 * len(C.AUTHORS[3:]) // 3)
        spacer = new_p("Author"); ppr_add(spacer, el("w:spacing", **{"w:before": 0, "w:after": 0, "w:line": 120,
                                                                      "w:lineRule": "exact"}))
        title_end.addprevious(spacer); title_end.addprevious(tB)
    # title_end paragraph keeps the first-page sectPr (1 column); tighten its spacing
    tpr = title_end.find(qn("w:pPr"))
    sp = tpr.find(qn("w:spacing"))
    for a in list(sp.attrib):
        del sp.attrib[a]
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "240"); sp.set(qn("w:line"), "120"); sp.set(qn("w:lineRule"), "exact")

    # ---- abstract & keywords
    p = insert(new_p("Abstract"))
    p.append(run("Abstract", bold=True, italic=True)); p.append(run("—", bold=True))
    rich(p, cite(C.ABSTRACT), bold=True)
    p = insert(new_p("Keywords"))
    p.append(run("Keywords—", bold=True, italic=True))
    rich(p, C.KEYWORDS, bold=True, italic=True)

    counters = {"eq": 0}

    def section_break(prev_p, cols):
        """Attach a continuous section break to prev_p; the section ending there gets `cols` columns.
        If prev_p already closes a single-column (full-width) section and a new full-width element follows,
        that break is removed instead, so consecutive full-width elements share one section."""
        ppr = prev_p.find(qn("w:pPr"))
        old = ppr.find(qn("w:sectPr")) if ppr is not None else None
        if old is not None:
            if cols == 2:
                ppr.remove(old)
                return
            raise ValueError("paragraph already carries a section break")
        s = copy.deepcopy(body_sect)
        c = s.find(qn("w:cols"))
        if cols == 1:
            for a in list(c.attrib):
                del c.attrib[a]
            c.set(qn("w:space"), "720")
        ppr_add(prev_p, s)

    last_p = [None]

    def body_para(text):
        p = insert(new_p("BodyText")); rich(p, cite(text)); last_p[0] = p; return p

    def equation(latex):
        counters["eq"] += 1
        p = insert(new_p("equation"))
        tabs = el("w:tabs")
        tabs.append(el("w:tab", **{"w:val": "center", "w:pos": COL_W // 2}))
        tabs.append(el("w:tab", **{"w:val": "right", "w:pos": COL_W}))
        ppr_add(p, tabs)
        ppr_add(p, el("w:spacing", **{"w:before": 60, "w:after": 60, "w:line": 240, "w:lineRule": "auto"}))
        p.append(run("\t", font=TNR))
        if PREVIEW["fn"]:
            p.append(picture_run(*PREVIEW["fn"](latex, 10)))
        else:
            p.append(sized_math(MATH[latex], 10))
        p.append(run(f"\t({counters['eq']})", font=TNR))
        last_p[0] = p

    def caption_fig(text):
        p = insert(new_p("figurecaption")); rich(p, cite(text)); return p

    def figure(path, width_in, caption, full=False):
        if full:
            section_break(last_p[0], 2)
        p = insert(new_p("Normal"))
        ppr_add(p, el("w:keepNext"))
        ppr_add(p, el("w:spacing", **{"w:before": 120, "w:after": 0}))
        ppr_add(p, el("w:jc", **{"w:val": "center"}))
        p.append(picture_run(path, width_in=width_in))
        c = caption_fig(caption)
        if full:
            ppr_add(c, el("w:jc", **{"w:val": "both"}))
            section_break(c, 1)
        last_p[0] = c

    def table(spec, rows, full=False, font=8):
        if full:
            section_break(last_p[0], 2)
        h = insert(new_p("tablehead")); h.append(run(spec["caption"]))
        ppr_add(h, el("w:keepNext"))
        widths = spec["widths"]; scale = (BODY_W if full else COL_W) / sum(widths) * 0.995
        wt = [int(w * scale) for w in widths]
        t = el("w:tbl"); tp = el("w:tblPr"); t.append(tp)
        tp.append(el("w:tblW", **{"w:w": sum(wt), "w:type": "dxa"}))
        tp.append(el("w:jc", **{"w:val": "center"}))
        b = el("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b.append(el(f"w:{e}", **{"w:val": "single", "w:sz": 2, "w:space": 0, "w:color": "auto"}))
        tp.append(b); tp.append(el("w:tblLayout", **{"w:type": "fixed"}))
        m = el("w:tblCellMar")
        for k in ("top", "left", "bottom", "right"):
            m.append(el(f"w:{k}", **{"w:w": 15 if k in ("top", "bottom") else 50, "w:type": "dxa"}))
        tp.append(m)
        g = el("w:tblGrid"); t.append(g)
        for w in wt:
            g.append(el("w:gridCol", **{"w:w": w}))
        allrows = [spec["header"]] + rows
        for ri, r in enumerate(allrows):
            tr = el("w:tr"); t.append(tr); trpr = el("w:trPr"); tr.append(trpr); trpr.append(el("w:cantSplit"))
            if ri == 0:
                trpr.append(el("w:tblHeader"))
            for ci, (val, w) in enumerate(zip(r, wt)):
                tc = el("w:tc"); tr.append(tc)
                tcp = el("w:tcPr"); tc.append(tcp); tcp.append(el("w:tcW", **{"w:w": w, "w:type": "dxa"}))
                tcp.append(el("w:vAlign", **{"w:val": "center"}))
                p = new_p("tablecolhead" if ri == 0 else "tablecopy"); tc.append(p)
                ppr_add(p, el("w:keepNext")) if (ri < len(allrows) - 1 or spec.get("note")) else None
                ppr_add(p, el("w:jc", **{"w:val": "center" if (ri == 0 or ci > 0) else "left"}))
                rich(p, cite(val), size=font, bold=(ri == 0))
        insert(t)
        after = insert(new_p("Normal"))
        if spec.get("note"):
            ppr_add(after, el("w:spacing", **{"w:before": 20, "w:after": 80}))
            ppr_add(after, el("w:jc", **{"w:val": "both"}))
            rich(after, cite(spec["note"]), size=7)
        else:
            ppr_add(after, el("w:spacing", **{"w:before": 0, "w:after": 60, "w:line": 120, "w:lineRule": "exact"}))
        last_p[0] = after
        if full:
            section_break(after, 1)

    FIGS = C.FIGS
    for h1, h2, paras in C.SECTIONS:
        if h1:
            p = insert(new_p("Heading1")); p.append(run(h1)); last_p[0] = p
        if h2:
            p = insert(new_p("Heading2")); p.append(run(h2)); last_p[0] = p
        for para in paras:
            if para.startswith("EQ:"):
                equation(para[3:])
            elif para.startswith("FIG:"):
                k = para[4:]; fpath, w, full = FIGS[k]; figure(fpath, w, C.CAPTIONS[k], full)
            elif para.startswith("TAB:"):
                k = para[4:]; T = C.TABLES[k]
                table(T, T["rows"], full=T.get("full", False), font=T.get("font", 7.5))
            else:
                p = body_para(para)
                # paragraphs that continue after an equation carry no first-line indent
                prev = p.getprevious()
                if prev is not None and prev.find(qn("w:pPr")) is not None:
                    st = prev.find(qn("w:pPr")).find(qn("w:pStyle"))
                    if st is not None and st.get(qn("w:val")) == "equation" and para[:1].islower():
                        ppr_add(p, el("w:ind", **{"w:firstLine": 0}))

    # ---- references
    p = insert(new_p("Heading5")); p.append(run("References"))
    for k in order:
        p = insert(new_p("references")); rich(p, REFS[k])
    unused = [k for k in REFS if k not in order]
    assert not unused, unused
    doc.save(path)
    return order


if __name__ == "__main__":
    collect_math()
    out = sys.argv[1] if len(sys.argv) > 1 else "out/PreOp_XAI_IEEE.docx"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    o = build(out)
    print("saved", out, "refs", len(o))
